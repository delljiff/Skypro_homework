import csv
import json
import os
from typing import Dict, List


def read_json(filepath: str) -> List[Dict]:
    """
    Читает JSON-файл с транзакциями.
    Возвращает список словарей с единой структурой.
    """
    if not os.path.exists(filepath):
        print(f"Ошибка: файл {filepath} не найден")
        return []

    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    transactions = []
    for item in data:
        if not item:  # пропускаем пустые словари
            continue

        # Для JSON сумма и валюта вложены в operationAmount
        amount = float(item.get("operationAmount", {}).get("amount", 0))
        currency = item.get("operationAmount", {}).get("currency", {}).get("code", "")

        transactions.append(
            {
                "id": item.get("id"),
                "state": item.get("state"),
                "date": item.get("date"),
                "amount": amount,
                "currency": currency,
                "description": item.get("description", ""),
                "from": item.get("from", ""),
                "to": item.get("to", ""),
            }
        )
    return transactions


def read_csv(filepath: str) -> List[Dict]:
    """
    Читает CSV-файл с транзакциями.
    Возвращает список словарей с единой структурой.
    """
    if not os.path.exists(filepath):
        print(f"Ошибка: файл {filepath} не найден")
        return []

    transactions = []
    with open(filepath, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter=";")
        for row in reader:
            if not row.get("id"):  # пропускаем пустые строки
                continue

            transactions.append(
                {
                    "id": int(row["id"]) if row.get("id") else None,
                    "state": row.get("state"),
                    "date": row.get("date"),
                    "amount": float(row["amount"]) if row.get("amount") else 0,
                    "currency": row.get("currency_code", ""),
                    "description": row.get("description", ""),
                    "from": row.get("from", ""),
                    "to": row.get("to", ""),
                }
            )
    return transactions


def read_xlsx(filepath: str) -> List[Dict]:
    """
    Читает XLSX-файл с транзакциями.
    Возвращает список словарей с единой структурой.
    """
    if not os.path.exists(filepath):
        print(f"Ошибка: файл {filepath} не найден")
        return []

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Для работы с XLSX установите библиотеку: pip install openpyxl")
        return []

    wb = load_workbook(filepath, data_only=True)
    ws = wb.active

    # Первая строка — заголовки
    headers = [cell.value for cell in ws[1]]

    transactions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:  # если id пустой — пропускаем
            continue

        row_dict = dict(zip(headers, row))
        transactions.append(
            {
                "id": int(row_dict["id"]) if row_dict.get("id") else None,
                "state": row_dict.get("state"),
                "date": row_dict.get("date"),
                "amount": float(row_dict["amount"]) if row_dict.get("amount") else 0,
                "currency": row_dict.get("currency_code", ""),
                "description": row_dict.get("description", ""),
                "from": row_dict.get("from", ""),
                "to": row_dict.get("to", ""),
            }
        )
    return transactions
